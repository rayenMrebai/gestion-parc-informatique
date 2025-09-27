import streamlit as st
from sqlalchemy import text
from database import SessionLocal
import mysql.connector
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors
import hashlib
from crud import add_project, update_project, delete_project, fetch_projects, project_exists
from crud import add_ligne, update_ligne, delete_ligne, fetch_lignes, ligne_exists
from crud import add_poste, update_poste, delete_poste, fetch_postes, poste_exists
from crud import add_equipement, update_equipement, delete_equipement, fetch_equipements, deplacer_equipement, equipement_existe, get_id_by_num_serie, get_emplacement, get_equipements, liste_equipements
from crud import add_historique_ac

st.set_page_config(page_title="Home", layout="wide")

if "authenticated" not in st.session_state or not st.session_state["authenticated"]:
    st.warning("Vous devez vous connecter d'abord!")
    st.switch_page("app.py")



def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()
db = next(get_db())

def get_filter_options():
    db = next(get_db())
    try:
        names = [row[0] for row in db.execute(
            text("SELECT DISTINCT host_name FROM equipements WHERE host_name IS NOT NULL")
        ).fetchall()]
        names.insert(0, 'Tous')

        modeles = [row[0] for row in db.execute(
            text("SELECT DISTINCT modele FROM equipements")
        ).fetchall()]

        types = [row[0] for row in db.execute(
            text("SELECT DISTINCT equipement FROM equipements")
        ).fetchall()]

        return names, modeles, types
    finally:
        db.close()



def filtrer_equipements(name, modele, type_):
    where_clauses = []
    params = {}

    if name == 'Tous':
        where_clauses.append("1=1")
    elif name == 'None':
        where_clauses.append("(e.host_name IS NULL OR e.host_name = '')")
    else:
        where_clauses.append("e.host_name LIKE :name")
        params['name'] = f"%{name}%"

    if modele == 'Tous':
        where_clauses.append("1=1")
    else:
        where_clauses.append("e.modele LIKE :modele")
        params['modele'] = f"%{modele}%"

    if type_ == 'Tous':
        where_clauses.append("1=1")
    else:
        where_clauses.append("e.equipement LIKE :type")
        params['type'] = f"%{type_}%"

    query = f"""
        SELECT e.host_name, e.modele, e.equipement, e.num_serie,
               p.name AS projet, l.name AS ligne, po.name AS poste
        FROM equipements e
        JOIN postes po ON e.poste_id = po.id
        JOIN lignes l ON po.ligne_id = l.id
        JOIN projects p ON l.project_id = p.id
        WHERE {' AND '.join(where_clauses)}
    """

    resultats = db.execute(text(query), params).mappings().all()
    return resultats

def excel(data):
    df = pd.DataFrame(data)
    output = BytesIO()
    writer = pd.ExcelWriter(output, engine='openpyxl')
    df.to_excel(writer, index=False, sheet_name='Équipements')
    workbook = writer.book
    worksheet = writer.sheets['Équipements']
    worksheet.insert_rows(1)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
    worksheet.cell(row=1, column=1).value = "Résultats"
    worksheet.cell(row=1, column=1).font = Font(bold=True, size=14)

    for col_num, column_title in enumerate(df.columns, 1):
        cell = worksheet.cell(row=2, column=col_num)
        cell.font = Font(bold=True)

    for i, col in enumerate(df.columns, 1):
        max_length = max(df[col].astype(str).map(len).max(), len(col))
        worksheet.column_dimensions[get_column_letter(i)].width = max_length + 2

    writer.close()
    output.seek(0)
    return output

def pdf(data):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    table_data = [["Host Nom", "Modèle", "Equipement", "N° Série", "Projet", "Ligne", "Poste"]]
    
    for row in data:
        host_name = row["host_name"] if row["host_name"] else "None"
        table_data.append([
            host_name,
            row["modele"],
            row["equipement"],
            row["num_serie"],
            row["projet"],
            row["ligne"],
            row["poste"]
        ])
    
    table = Table(table_data)
    style = TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.black), 
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), 
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ])
    table.setStyle(style)
    doc.build([table])
    buffer.seek(0)
    return buffer


def import_excel(uploaded_file, db):
    df_raw = pd.read_excel(uploaded_file, header=None)
    if df_raw.empty or df_raw.shape[0] < 2:
        return ["Fichier vide ou insuffisant"]

    def detect_header_row(df):
        for i, row in df.iterrows():
            if 'Projet' in row.values and 'Ligne' in row.values:
                return i
        return 0

    header_row = detect_header_row(df_raw)
    df = pd.read_excel(uploaded_file, skiprows=header_row)
    df.columns = df.columns.str.strip()

    required_cols = ['Projet', 'Ligne', 'Poste', 'Equipement', 'Modele', 'Numero Serie','Modele']
    if not all(col in df.columns for col in required_cols):
        return [f"Colonnes attendues absentes : {', '.join(required_cols)}"]

    df = df.dropna(subset=['Projet', 'Ligne', 'Poste', 'Numero Serie', 'Equipement', 'Modele'])

    erreurs = set()

    for i, row in df.iterrows():
        projet = str(row['Projet']).strip()
        ligne_nom = str(row['Ligne']).strip()
        poste = str(row['Poste']).strip()
        equipement = str(row['Equipement']).strip()
        modele = str(row['Modele']).strip()
        num_serie = str(row['Numero Serie']).strip()
        host_name = str(row.get('Host name', '')).strip()

        champs = [projet, ligne_nom, poste, equipement, modele, num_serie]
        if any(c == '' or c.lower() == 'none' for c in champs):
            erreurs.add(f"Ligne ignorée pour champs obligatoires vides ou 'None' : Projet='{projet}', Ligne='{ligne_nom}', Poste='{poste}', Equipement='{equipement}', Modele='{modele}', Num Serie='{num_serie}'")
            continue

        try:
            result = db.execute(text("SELECT id FROM projects WHERE name = :name"), {"name": projet}).fetchone()
            if result:
                project_id = result[0]
            else:
                db.execute(text("INSERT INTO projects (name) VALUES (:name)"), {"name": projet})
                db.commit()
                project_id = db.execute(text("SELECT LAST_INSERT_ID()")).scalar()

            result = db.execute(text("SELECT id FROM lignes WHERE name = :name"), {"name": ligne_nom}).fetchone()
            if result:
                ligne_id = result[0]
            else:
                db.execute(text("INSERT INTO lignes (name, project_id) VALUES (:name, :project_id)"),{"name": ligne_nom, "project_id": project_id})
                db.commit()
                ligne_id = db.execute(text("SELECT LAST_INSERT_ID()")).scalar()

            result = db.execute(text("SELECT id FROM postes WHERE name = :name"), {"name": poste}).fetchone()
            if result:
                poste_id = result[0]
            else:
                db.execute(text("INSERT INTO postes (name, ligne_id) VALUES (:name, :ligne_id)"),{"name": poste, "ligne_id": ligne_id})
                db.commit()
                poste_id = db.execute(text("SELECT LAST_INSERT_ID()")).scalar()

            result = db.execute(text("SELECT id FROM equipements WHERE num_serie = :num_serie"), {"num_serie": num_serie}).fetchone()
            if result:
                erreurs.add("Numéro de série déjà existant ignoré.")
                continue
            else:
                db.execute(
                    text("INSERT INTO equipements (poste_id, host_name, equipement, modele, num_serie) VALUES (:poste_id, :host_name, :equipement, :modele, :num_serie)"),
                    {"poste_id": poste_id, "host_name": host_name or None, "equipement": equipement, "modele": modele, "num_serie": num_serie}
                )
                db.commit()

        except Exception as e:
            db.rollback()
            erreurs.add(f"Erreur à la ligne avec projet '{projet}', ligne '{ligne_nom}', poste '{poste}': {e}")

    return list(erreurs)


def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def verifier_password(input_password):
    return hash_password(input_password) == st.session_state["password"]

def action(db, action, cible, details=""):
    add_historique_ac(
        db=db,
        utilisateur_id=st.session_state.get("id"),
        nom=st.session_state.get("nom"),
        role=st.session_state.get("role"),
        action=action,
        cible=cible,
        details=details
    )

def reset_editing():
    st.session_state.pop('editing_project', None)
    st.session_state.pop('editing_ligne', None)
    st.session_state.pop('editing_poste', None)
    st.session_state.pop('editing_equipement', None)

def show_projects():
    st.session_state.page = 'projects'
    st.session_state.selected_project = None
    reset_editing()

def show_lignes(project_id):
    st.session_state.page = 'lignes'
    st.session_state.selected_project = project_id
    reset_editing()

def show_postes(ligne_id):
    st.session_state.page = 'postes'
    st.session_state.selected_ligne = ligne_id
    reset_editing()

def show_equipements(poste_id):
    st.session_state.page = 'equipements'
    st.session_state.selected_poste = poste_id
    reset_editing()




if 'page' not in st.session_state:
    st.session_state.page = 'projects'
if 'selected_project' not in st.session_state:
    st.session_state.selected_project = None
if 'selected_ligne' not in st.session_state:
    st.session_state.selected_ligne = None
if 'selected_poste' not in st.session_state:
    st.session_state.selected_poste = None


if "show_filter" not in st.session_state:
    st.session_state.show_filter = False

if "show_delete" not in st.session_state:
    st.session_state.show_delete = None

st.title(f"Session {st.session_state['role']}!")

if st.session_state.page == 'projects':
    col_import, col_filter = st.columns(2)
    if st.session_state['role'] == 'admin':
        with col_import:
            uploaded_file = st.file_uploader("Importer un fichier Excel", type=["xlsx"])
            if uploaded_file:
                if st.button("Importer"):
                    db = next(get_db())
                    try:
                        erreurs = import_excel(uploaded_file, db)
                        if erreurs:
                            st.warning("Données ignorées pour les raisons suivantes :")
                            for err in erreurs:
                                st.text(err)
                        else:
                            st.success("Importation réussie sans conflits !")
                    finally:
                        db.close()


    with col_filter:
        if st.button("Filtrer", key="filtrer_button", use_container_width=True):
            st.session_state.show_filter = True

        names, modeles, types = get_filter_options()

        if st.session_state.get("show_filter", False):
            selected_name = st.selectbox("Host name", ["Tous", "None"] + names)
            selected_modele = st.selectbox("Modèle", ["Tous"] + modeles)
            selected_type = st.selectbox("Equipement", ["Tous"] + types)

            if st.button("Appliquer les filtres", key="appliquer_button"):
                resultats = filtrer_equipements(selected_name, selected_modele, selected_type)
                if resultats:
                    df = pd.DataFrame(resultats)
                    st.dataframe(df)
                    st.write(f"**Total : {len(resultats)} équipements trouvés**")
                    excel_file = excel(resultats)
                    st.download_button(
                        label="Télécharger en Excel",
                        data=excel_file,
                        file_name="equipements_filtrés.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    pdf_file = pdf(resultats)
                    st.download_button(
                        label="Télécharger en PDF",
                        data=pdf_file,
                        file_name="equipements_filtrés.pdf",
                        mime="application/pdf"
                    )
                else:
                    st.warning("Aucun équipement trouvé.")
            if st.button("Annuler", key="annuler_button"):
                st.session_state.show_filter = False
                st.rerun()
    search_term = st.text_input("", placeholder="Rechercher un équipement")
    if search_term:
        equipements = get_equipements(db,search_term)
        if equipements:
            df = pd.DataFrame(equipements)
            st.dataframe(df)
            st.markdown(f"**Nombre total : {len(df)} équipements trouvés.**")
            excel_file = excel(equipements)
            st.download_button(
                label="Télécharger en Excel",
                data=excel_file,
                file_name="equipements_recherchés.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            pdf_file = pdf(equipements)
            st.download_button(
                label="Télécharger en PDF",
                data=pdf_file,
                file_name="equipements_filtrés.pdf",
                mime="application/pdf"
            )
        else:
            st.warning("Aucun équipement trouvé.")
    st.header("Gestion des Projets")

    if st.session_state['role'] == 'admin':
        with st.form("Nouveau Projet"):
            new_project = st.text_input("Nom du nouveau projet")
            if st.form_submit_button("Créer Projet"):
                if new_project:
                    if project_exists(db,new_project):
                        st.error("Le nom du projet existe déjà.")
                    else:
                        add_project(db,new_project)
                        action(db, "création", "projet", f"Projet '{new_project}' ajouté")
                        st.success("Projet ajouté avec succès.")
                        st.rerun()
                else:
                    st.error("Veuillez remplir le champ.")

    projects = fetch_projects(db)
    for project in projects:
        with st.container():
            col1, col2, col3 = st.columns([6, 2, 2])
            with col1:
                if st.button(project.name, key=f"project_{project.id}", use_container_width=True):
                    show_lignes(project.id)
            if st.session_state['role'] == 'admin':
                with col2:
                    if st.button("Modifier", key=f"edit_project_{project.id}", use_container_width=True):
                        st.session_state.editing_project = project.id
                with col3:
                    if st.button("Supprimer", key=f"del_project_{project.id}", use_container_width=True):
                        st.session_state.show_delete = project.id
                if st.session_state.get('editing_project') == project.id:
                    new_name = st.text_input("Nouveau nom", value=project.name, key=f"edit_input_{project.id}")
                    if st.button("Valider", key=f"validate_{project.id}"):
                        if new_name:
                            if project_exists(db,new_name):
                                st.error("Un projet avec ce nom existe déjà.")
                            else:
                                old_name= project.name
                                update_project(db,project.id, new_name)
                                action(db, "modification", "projet", f"Projet renommé de '{old_name}' en '{new_name}'")
                                reset_editing()
                                st.rerun()
                        else:
                            st.error("Veuillez remplir le champ.")
                if st.session_state.show_delete == project.id:
                    with st.form(f"confirm_delete_form_{project.id}"):
                        input_password= st.text_input("Mot de passe:",type="password")
                        confirm= st.form_submit_button("Valider")
                        if confirm:
                            if verifier_password(input_password):
                                delete_project(db,project.id)
                                action(db, "suppression", "projet", f"Projet {project.name} supprimé")
                            else: 
                                st.error("mot de passe incorrect.")
                            st.session_state.show_delete = None
                            st.rerun()

elif st.session_state.page == 'lignes':
    st.button("← Retour aux Projets", on_click=show_projects)
    st.header("Gestion des Lignes")

    if st.session_state['role'] == 'admin':
        with st.form("Nouvelle Ligne"):
            new_ligne = st.text_input("Nom de la nouvelle ligne")
            if st.form_submit_button("Créer Ligne"):
                if new_ligne:
                    if ligne_exists(db,new_ligne):
                        st.error("Une ligne avec ce nom existe déjà.")
                    else:
                        add_ligne(db,st.session_state.selected_project, new_ligne)
                        action(db, "création", "ligne", f"Ligne '{new_ligne}' ajoutée")
                        st.rerun()
                else:
                    st.error("Veuillez remplir champ.")

    lignes = fetch_lignes(db,st.session_state.selected_project)
    for ligne in lignes:
        with st.container():
            col1, col2, col3 = st.columns([6, 2, 2])
            with col1:
                if st.button(ligne.name, key=f"ligne_{ligne.id}", use_container_width=True):
                    show_postes(ligne.id)
            if st.session_state['role'] == 'admin':
                with col2:
                    if st.button("Modifier", key=f"edit_ligne_{ligne.id}", use_container_width=True):
                        st.session_state.editing_ligne = ligne.id
                with col3:
                    if st.button("Supprimer", key=f"del_ligne_{ligne.id}", use_container_width=True):
                        st.session_state.show_delete = ligne.id
                if st.session_state.get('editing_ligne') == ligne.id:
                    new_name = st.text_input("Nouveau nom", value=ligne.name, key=f"edit_input_{ligne.id}")
                    if st.button("Valider", key=f"validate_{ligne.id}"):
                        if ligne_exists(db,new_name):
                            st.error("Une ligne avec ce nom existe déjà.")
                        else:
                            old_name=ligne.name
                            update_ligne(db,ligne.id, new_name)
                            action(db, "modification", "ligne", f"Ligne renommé de '{old_name}' en '{new_name}'")
                            reset_editing()
                            st.rerun()
                if st.session_state.show_delete == ligne.id:
                    with st.form(f"confirm_delete_form_{ligne.id}"):
                        input_password= st.text_input("Mot de passe:",type="password")
                        confirm= st.form_submit_button("Valider")
                        if confirm:
                            if verifier_password(input_password):
                                delete_ligne(db,ligne.id)
                                action(db, "suppression", "ligne", f"Ligne {ligne.name} supprimée")
                            else: 
                                st.error("mot de passe incorrect.")
                            st.session_state.show_delete = None
                            st.rerun()

elif st.session_state.page == 'postes':
    st.button("← Retour aux Lignes", on_click=lambda: show_lignes(st.session_state.selected_project))
    st.header("Gestion des Postes")

    if st.session_state['role'] == 'admin':
        with st.form("Nouveau Poste"):
            new_poste = st.text_input("Nom du nouveau poste")
            if st.form_submit_button("Créer Poste"):
                if new_poste:
                    if poste_exists(db,new_poste):
                        st.error("Un poste avec ce nom existe déjà.")
                    else:
                        add_poste(db,st.session_state.selected_ligne, new_poste)
                        action(db, "création", "poste", f"Poste '{new_poste}' ajouté")
                        st.rerun()
                else:
                    st.error("Veuillez remplir champ.")

    postes = fetch_postes(db,st.session_state.selected_ligne)
    for poste in postes:
        with st.container():
            col1, col2, col3 = st.columns([6, 2, 2])
            with col1:
                if st.button(poste.name, key=f"poste_{poste.id}", use_container_width=True):
                    show_equipements(poste.id)
            if st.session_state['role'] == 'admin':
                with col2:
                    if st.button("Modifier", key=f"edit_poste_{poste.id}", use_container_width=True):
                        st.session_state.editing_poste = poste.id
                with col3:
                    if st.button("Supprimer", key=f"del_poste_{poste.id}", use_container_width=True):
                        st.session_state.show_delete = poste.id
                if st.session_state.get('editing_poste') == poste.id:
                    new_name = st.text_input("Nouveau nom", value=poste.name, key=f"edit_input_{poste.id}")
                    if st.button("Valider", key=f"validate_{poste.id}"):
                        if poste_exists(db,new_name):
                            st.error("Un poste avec ce nom existe déjà.")
                        else:
                            old_name=poste.name
                            update_poste(db,poste.id, new_name)
                            action(db, "modification", "poste", f"Poste '{old_name}' renommé en '{new_name}'")
                            reset_editing()
                            st.rerun()
                if st.session_state.show_delete == poste.id:
                    with st.form(f"confirm_delete_form_{poste.id}"):
                        input_password= st.text_input("Mot de passe:",type="password")
                        confirm= st.form_submit_button("Valider")
                        if confirm:
                            if verifier_password(input_password):
                                delete_poste(db,poste.id)
                                action(db, "suppression", "poste", f"Poste '{poste.name}' supprimé")
                            else: 
                                st.error("mot de passe incorrect.")
                            st.session_state.show_delete = None
                            st.rerun()

elif st.session_state.page == 'equipements':
    st.button("← Retour aux Postes", on_click=lambda: show_postes(st.session_state.selected_ligne))
    st.header("Gestion des Équipements")

    if st.session_state['role'] == 'admin':
        scanned_code = st.text_input("Scannez un code-barres ici")
        if scanned_code:
            if equipement_existe(db,scanned_code):
                projet, ligne, poste = get_emplacement(db,scanned_code)
                st.success(f"Équipement trouvé : Projet {projet}, Ligne {ligne}, Poste {poste}")
            else:
                st.warning("Équipement non trouvé. Veuillez saisir les informations manuellement.")
                with st.form("ajouter_eq"):
                    name = st.text_input("Host name")
                    modele = st.text_input("Modèle")
                    type_ = st.text_input("Equipement")
                    st.text_input("Numéro de série", value=scanned_code, disabled=True)
                    submitted = st.form_submit_button("Ajouter")
                    if submitted:
                        add_equipement(db,st.session_state.selected_poste, name, type_, modele, scanned_code)
                        action(db, "création", "équipement", f"Équipement avec numéro de série '{scanned_code}' ajouté")
                        st.success("Équipement ajouté.")
                        st.rerun()
        if 'show_add_form' not in st.session_state:
            st.session_state['show_add_form'] = False
        
        if 'deplacement' not in st.session_state:
            st.session_state['deplacement'] = False
        col_add, col_deplacer = st.columns(2)
        with col_add:
            if st.button("Ajouter un équipement", use_container_width=True):
                st.session_state['show_add_form'] = not st.session_state['show_add_form']

            if st.session_state['show_add_form']:
                with st.form("Nouvel Équipement"):
                    name = st.text_input("Host name")
                    type_ = st.text_input("Equipement")
                    modele = st.text_input("Modèle")
                    num_serie = st.text_input("Numéro de Série")

                    if st.form_submit_button("Ajouter Équipement"):
                        if type_ and modele and num_serie:
                            if equipement_existe(db,num_serie):
                                st.error("Un équipement avec ce numéro de série existe déjà.")
                            else:
                                add_equipement(db,st.session_state.selected_poste, name, type_, modele, num_serie)
                                action(db, "création", "équipement", f"Équipement avec numéro de série '{num_serie}' ajouté")
                                st.session_state['show_add_form'] = False
                                st.rerun()
                        else:
                            st.error("Veuillez remplir tous les champs.")
        with col_deplacer:
            if st.button("Déplacer", use_container_width=True):
                st.session_state['deplacement'] = not st.session_state['deplacement']
            if st.session_state['deplacement']:
                num_serie = st.text_input("Numéro de série de l'équipement")
                if num_serie:
                    if equipement_existe(db,num_serie):
                        eq_id = get_id_by_num_serie(db,num_serie)
                        old_emplacement=get_emplacement(db,num_serie)
                        projets = fetch_projects(db)  
                        projet = st.selectbox("Choisir un projet", projets, format_func=lambda x: x.name)

                        if projet:
                            lignes = fetch_lignes(db,projet.id)
                            ligne = st.selectbox("Choisir une ligne", lignes, format_func=lambda x: x.name)

                            if ligne:
                                postes = fetch_postes(db,ligne.id)
                                poste = st.selectbox("Choisir un poste", postes, format_func=lambda x: x.name)

                                if poste:
                                    if st.button("Déplacer"):
                                        deplacer_equipement(db,eq_id, poste.id)
                                        nouveau_emplacement = f"{projet.name} / {ligne.name} / {poste.name}"
                                        action(db, "déplacement", "équipement", f"Équipement numéro de série '{num_serie}' déplacé de '{old_emplacement}' à '{nouveau_emplacement}'")
                                        st.session_state['deplacement']= False
                                        st.rerun()
                    else:
                        st.error("Équipement non trouvé.")
                else:
                    st.error("Veuillez remplir cahmp.")


    equipements = fetch_equipements(db,st.session_state.selected_poste)
    equi= liste_equipements(db,st.session_state.selected_poste)
    if equipements:
        col1, col2 = st.columns(2)
        with col1:
            excel_file = excel(equi)
            st.download_button(
                label="Exporter Excel",
                data=excel_file,
                file_name="equipements_emplacement.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with col2:
            pdf_file = pdf(equi)
            st.download_button(
                label="Exporter PDF",
                data=pdf_file,
                file_name="equipements_emplacement.pdf",
                mime="application/pdf",
                use_container_width=True
            )

    if equipements:
        st.markdown("Liste des équipements")
        header_cols = st.columns([2, 2, 2, 2, 1, 1])
        with header_cols[0]:
            st.markdown("Host name")
        with header_cols[1]:
            st.markdown("Équipement")
        with header_cols[2]:
            st.markdown("Modèle")
        with header_cols[3]:
            st.markdown("Numéro de série")
        with header_cols[4]:
            st.markdown("Modifier")
        with header_cols[5]:
            st.markdown("Supprimer")

        for equipement in equipements:
            eq_id = equipement.id
            name = equipement.host_name or "None"
            type_ = equipement.equipement
            modele = equipement.modele
            num_serie = equipement.num_serie
            
            row_cols = st.columns([2, 2, 2, 2, 1, 1])
            
            with row_cols[0]:
                st.text(name)
            with row_cols[1]:
                st.text(type_)
            with row_cols[2]:
                st.text(modele)
            with row_cols[3]:
                st.text(num_serie)

            if st.session_state['role'] == 'admin':
                with row_cols[4]:
                    if st.button("Modifier", key=f"edit_eq_{eq_id}", use_container_width=True):
                        st.session_state.editing_equipement = eq_id
                with row_cols[5]:
                    if st.button("Supprimer", key=f"del_eq_{eq_id}", use_container_width=True):
                        st.session_state.show_delete = eq_id

            if st.session_state.get('editing_equipement') == eq_id:
                st.markdown("Modifier l'équipement")
                new_name = st.text_input("Host name", value=name, key=f"edit_eq_name_{eq_id}")
                new_type = st.text_input("Equipement", value=type_, key=f"edit_eq_type_{eq_id}")
                new_modele = st.text_input("Modèle", value=modele, key=f"edit_eq_modele_{eq_id}")
                new_num_serie = st.text_input("Numéro de Série", value=num_serie, key=f"edit_eq_sn_{eq_id}")

                if st.button("Valider", key=f"validate_eq_{eq_id}"):
                    if new_type and new_modele and new_num_serie:
                        if equipement_existe(db,new_num_serie) and new_num_serie != num_serie:
                            st.error("Un autre équipement existe déjà avec ce numéro de série.")
                        else:
                            old_num_serie=num_serie
                            update_equipement(db,eq_id, new_name, new_type, new_modele, new_num_serie)
                            action(db, "modification", "équipement", f"Équipement numéro de série '{old_num_serie}' modifié en '{new_num_serie}'")
                            reset_editing()
                            st.rerun()
                    else:
                        st.error("Veuillez remplir tous les champs.")
            if st.session_state.show_delete == eq_id:
                with st.form(f"confirm_delete_form_{eq_id}"):
                    input_password= st.text_input("Mot de passe:",type="password")
                    confirm= st.form_submit_button("Valider")
                    if confirm:
                        if verifier_password(input_password):
                            delete_equipement(db,eq_id)
                            action(db, "suppression", "équipement", f"Équipement avec numéro de série '{num_serie}' supprimé")
                        else: 
                            st.error("mot de passe incorrect.")
                            st.session_state.show_delete = None
                            st.rerun()
    else:
        st.info("Aucun équipement trouvé pour ce poste.")



role = st.session_state.get("role", "user")

col = st.sidebar.container()

if col.button("Accueil", use_container_width=True):
    st.switch_page("pages/home.py")

if role == "admin":
    if col.button("Utilisateurs", use_container_width=True):
        st.switch_page("pages/users.py")
    if col.button("Historique Connexions", use_container_width=True):
        st.switch_page("pages/historique.py")
    if col.button("Historique Actions", use_container_width=True):
        st.switch_page("pages/actions.py")

if col.button("Déconnexion", use_container_width=True):
    st.session_state.clear()
    st.rerun()


